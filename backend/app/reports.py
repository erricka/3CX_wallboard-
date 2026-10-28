"""Daily report built from the stats stored in Postgres, exported as Excel or CSV."""

import csv
import io
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from .models import Agent, AgentDailyStats, CallRecord, OverallDailyStats, Queue
from .timeutil import PERIODS

COUNT_FIELDS = ("inbound", "outbound", "missed", "unanswered_slow")
TIME_FIELDS = ("talk_seconds", "work_seconds")
OVERALL_FIELDS = ("total", "inbound", "outbound", "internal", "inbound_answered", "inbound_missed",
                  "outbound_answered", "internal_answered", "talk_seconds", "wait_seconds", "max_wait_seconds")
OVERALL_HEADERS = ["Total calls", "Inbound calls", "Outbound calls", "Internal calls", "Inbound answered",
                   "Inbound missed", "Answer rate", "Outbound answered", "Talk time", "Average talk time",
                   "Average wait time", "Longest wait time"]


@dataclass
class ReportData:
    period: str
    days: list[date]
    tz: ZoneInfo
    generated_at: datetime
    slow_seconds: int
    agents: list[tuple[str, str]] = field(default_factory=list)  # (extension, name)
    queues: list[tuple[str, str]] = field(default_factory=list)  # (number, name): their members are the agents
    agent_daily: dict[tuple[date, str], dict] = field(default_factory=dict)
    overall_daily: dict[date, dict] = field(default_factory=dict)
    calls: list[CallRecord] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def title(self) -> str:
        first, last = self.days[0], self.days[-1]
        span = first.isoformat() if first == last else f"{first.isoformat()} to {last.isoformat()}"
        return f"{PERIODS[self.period]} ({span})"

    @property
    def filename_stem(self) -> str:
        return f"3cx-daily-report_{self.days[0].isoformat()}_to_{self.days[-1].isoformat()}"

    def agent_day(self, day: date, ext: str) -> dict:
        return self.agent_daily.get((day, ext)) or {f: 0 for f in COUNT_FIELDS + TIME_FIELDS}

    def overall_day(self, day: date) -> dict:
        return self.overall_daily.get(day) or {f: 0 for f in OVERALL_FIELDS}

    def overall_total(self) -> dict:
        days = [self.overall_day(d) for d in self.days]
        total = _totals(days, OVERALL_FIELDS)
        total["max_wait_seconds"] = max((d["max_wait_seconds"] for d in days), default=0)
        return total


async def load_report(session: AsyncSession, period: str, days: list[date], tz: ZoneInfo, queue_filter: list[str],
                      slow_seconds: int, now: datetime) -> ReportData:
    data = ReportData(period=period, days=days, tz=tz, generated_at=now, slow_seconds=slow_seconds)

    queue_stmt = select(Queue).order_by(Queue.number)
    if queue_filter:
        queue_stmt = queue_stmt.where(Queue.number.in_(queue_filter))
    queues = (await session.execute(queue_stmt)).scalars().all()
    data.queues = [(q.number, q.name) for q in queues]

    names = dict((await session.execute(select(Agent.extension, Agent.name))).all())
    agent_exts = set().union(*(q.agent_extensions or [] for q in queues)) if queues else set(names)

    for row in (await session.execute(select(AgentDailyStats).where(
            AgentDailyStats.day.in_(days), AgentDailyStats.extension.in_(agent_exts)))).scalars():
        data.agent_daily[(row.day, row.extension)] = {f: getattr(row, f) for f in COUNT_FIELDS + TIME_FIELDS}
        names.setdefault(row.extension, row.name)
    data.agents = sorted(((e, names.get(e) or e) for e in agent_exts), key=lambda a: (a[1].casefold(), a[0]))

    for row in (await session.execute(select(OverallDailyStats).where(OverallDailyStats.day.in_(days)))).scalars():
        data.overall_daily[row.day] = {f: getattr(row, f) for f in OVERALL_FIELDS}

    # Call log: main rows of every call on the PBX, the calls the overall counts are made of.
    by_call: dict[str, list[CallRecord]] = defaultdict(list)
    for rec in (await session.execute(select(CallRecord).where(CallRecord.day.in_(days))
                                      .order_by(CallRecord.start_time, CallRecord.segment_id))).scalars():
        by_call[rec.call_id or f"segment-{rec.segment_id}"].append(rec)
    for rows in by_call.values():
        main = [r for r in rows if r.indent == 0]
        data.calls.extend(main or rows)  # this PBX reports every row with Indent 1
    data.calls.sort(key=lambda r: (r.start_time or now, r.segment_id or 0))
    return data


def _totals(rows: list[dict], fields) -> dict:
    return {f: sum(r[f] for r in rows) for f in fields}


def _rate(answered: int, received: int) -> float | None:
    return answered / received if received else None


def overall_values(o: dict) -> list:
    """The OVERALL_HEADERS columns; times as seconds."""
    answered = o["inbound_answered"] + o["outbound_answered"] + o["internal_answered"]
    return [o["total"], o["inbound"], o["outbound"], o["internal"], o["inbound_answered"], o["inbound_missed"],
            _rate(o["inbound_answered"], o["inbound_answered"] + o["inbound_missed"]), o["outbound_answered"],
            o["talk_seconds"], o["talk_seconds"] / answered if answered else 0,
            o["wait_seconds"] / o["inbound_answered"] if o["inbound_answered"] else 0, o["max_wait_seconds"]]


OVERALL_TIME_COLUMNS = (8, 9, 10, 11)  # 0-based positions of the time values in overall_values
OVERALL_RATE_COLUMN = 6


def _hms(seconds: float) -> str:
    seconds = int(round(seconds))
    return f"{seconds // 3600:02d}:{seconds % 3600 // 60:02d}:{seconds % 60:02d}"


# ---- Excel ---------------------------------------------------------------------

HEADER_FILL = PatternFill("solid", fgColor="D9E1EC")
THIN = Side(style="thin", color="A6B1C2")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TIME_FORMAT = "[h]:mm:ss"


def _write_table(ws, start_row: int, headers: list[str], rows: list[list], formats: dict[int, str] | None = None) -> int:
    """Write a bordered table; returns the next free row."""
    for col, title in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col, value=title)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r, values in enumerate(rows, start_row + 1):
        for col, value in enumerate(values, 1):
            cell = ws.cell(row=r, column=col, value=value)
            cell.border = BORDER
            if formats and col in formats:
                cell.number_format = formats[col]
            if col > 1 and not isinstance(value, str):
                cell.alignment = Alignment(horizontal="center")
    return start_row + len(rows) + 1


def _autosize(ws, minimum: int = 10, maximum: int = 40) -> None:
    widths: dict[int, int] = defaultdict(int)
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None and not (cell.row <= 3 and cell.column == 1):
                widths[cell.column] = max(widths[cell.column], len(str(cell.value)))
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = min(max(width + 2, minimum), maximum)


def _overall_cells(o: dict) -> list:
    return [v / 86400 if i in OVERALL_TIME_COLUMNS else v for i, v in enumerate(overall_values(o))]


def _overall_formats(first_col: int) -> dict[int, str]:
    formats = {first_col + i: TIME_FORMAT for i in OVERALL_TIME_COLUMNS}
    formats[first_col + OVERALL_RATE_COLUMN] = "0.0%"
    return formats


def to_xlsx(data: ReportData) -> bytes:
    wb = Workbook()
    slow = f"Unanswered within {data.slow_seconds} sec"
    agent_headers = ["Inbound calls", "Outbound calls", "Missed calls", slow, "Talk time", "Work time"]

    # Summary
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "3CX Call Center Daily Report"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Period: {data.title}"
    ws["A3"] = f"Generated: {data.generated_at.astimezone(data.tz):%Y-%m-%d %H:%M} ({data.tz})"
    row = 5
    ws.cell(row=row, column=1, value="All calls on the PBX").font = Font(bold=True, size=12)
    row = _write_table(ws, row + 1, ["Period", *OVERALL_HEADERS], [[data.title, *_overall_cells(data.overall_total())]],
                       _overall_formats(2)) + 1
    ws.cell(row=row, column=1, value="Agents").font = Font(bold=True, size=12)
    agent_rows = []
    for ext, name in data.agents:
        t = _totals([data.agent_day(d, ext) for d in data.days], COUNT_FIELDS + TIME_FIELDS)
        agent_rows.append([name, ext, *(t[f] for f in COUNT_FIELDS), t["talk_seconds"] / 86400, t["work_seconds"] / 86400])
    row = _write_table(ws, row + 1, ["Agent", "Extension", *agent_headers], agent_rows,
                       {7: TIME_FORMAT, 8: TIME_FORMAT}) + 1
    if data.warnings:
        ws.cell(row=row, column=1, value="Notes").font = Font(bold=True, size=12)
        for i, warning in enumerate(data.warnings, row + 1):
            ws.cell(row=i, column=1, value=warning)
    _autosize(ws)

    # Daily agent stats (days where the agent worked or had calls)
    ws = wb.create_sheet("Daily Agents")
    rows = []
    for day in data.days:
        for ext, name in data.agents:
            s = data.agent_day(day, ext)
            if any(s[f] for f in COUNT_FIELDS + TIME_FIELDS):
                rows.append([day, name, ext, *(s[f] for f in COUNT_FIELDS), s["talk_seconds"] / 86400,
                             s["work_seconds"] / 86400])
    _write_table(ws, 1, ["Date", "Agent", "Extension", *agent_headers], rows,
                 {1: "yyyy-mm-dd", 8: TIME_FORMAT, 9: TIME_FORMAT})
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:I{len(rows) + 1}"
    _autosize(ws)

    # Daily totals of every call on the PBX
    ws = wb.create_sheet("Daily Overall")
    rows = [[day, *_overall_cells(data.overall_day(day))] for day in data.days]
    _write_table(ws, 1, ["Date", *OVERALL_HEADERS], rows, {1: "yyyy-mm-dd", **_overall_formats(2)})
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(OVERALL_HEADERS) + 1)}{len(rows) + 1}"
    _autosize(ws)

    # Call log
    ws = wb.create_sheet("Call Log")
    rows = []
    for c in data.calls:
        start = c.start_time.astimezone(data.tz).replace(tzinfo=None) if c.start_time else None
        rows.append([start, c.call_id, (c.raw or {}).get("Direction") or "", c.source_dn, c.source_name,
                     c.destination_dn, c.destination_name, c.action_name or c.action_dn,
                     "Yes" if c.answered else "No", c.ringing_seconds / 86400, c.talking_seconds / 86400, c.reason])
    _write_table(ws, 1, ["Start time", "Call ID", "Direction", "From", "From name", "To", "To name", "Via",
                         "Answered", "Ringing", "Talking", "Reason"], rows,
                 {1: "yyyy-mm-dd hh:mm:ss", 10: TIME_FORMAT, 11: TIME_FORMAT})
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:L{len(rows) + 1}"
    _autosize(ws)

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ---- CSV -----------------------------------------------------------------------

def to_csv(data: ReportData) -> bytes:
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(["3CX Call Center Daily Report", data.title])
    w.writerow([])
    w.writerow(["All calls on the PBX"])
    w.writerow(["Date", *OVERALL_HEADERS])
    for day in data.days:
        values = overall_values(data.overall_day(day))
        w.writerow([day.isoformat(), *(
            _hms(v) if i in OVERALL_TIME_COLUMNS else ("" if v is None else f"{v:.1%}") if i == OVERALL_RATE_COLUMN else v
            for i, v in enumerate(values))])
    w.writerow([])
    w.writerow(["Date", "Agent", "Extension", "Inbound calls", "Outbound calls", "Missed calls",
                f"Unanswered within {data.slow_seconds} sec", "Talk time", "Work time"])
    for day in data.days:
        for ext, name in data.agents:
            s = data.agent_day(day, ext)
            w.writerow([day.isoformat(), name, ext, *(s[f] for f in COUNT_FIELDS),
                        _hms(s["talk_seconds"]), _hms(s["work_seconds"])])
    for warning in data.warnings:
        w.writerow(["Note", warning])
    # BOM so Excel opens UTF-8 (Khmer names) correctly.
    return out.getvalue().encode("utf-8-sig")
